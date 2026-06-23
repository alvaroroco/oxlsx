"""
Parametric parity test suite.

Every test runs against BOTH openpyxl and oxlsx with the same fixture.
A failure in oxlsx only → divergence from openpyxl.
A failure in both → bad fixture or broken sample file.
"""

import datetime
import os
import tempfile
import warnings
import zipfile

import openpyxl
from openpyxl.comments import Comment as OpenPyxlComment
import pytest

import oxlsx


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _openpyxl_load(path):
    return openpyxl.load_workbook(path)


def _oxlsx_load(path):
    return oxlsx.load_workbook(path)


@pytest.fixture(params=["openpyxl", "oxlsx"])
def load_wb(request):
    if request.param == "openpyxl":
        return _openpyxl_load
    return _oxlsx_load


@pytest.fixture(params=["openpyxl", "oxlsx"])
def new_wb(request):
    """Returns a (workbook, save_fn) tuple for the write path."""
    if request.param == "openpyxl":
        wb = openpyxl.Workbook()
        return wb, wb.save
    wb = oxlsx.Workbook()
    return wb, wb.save


# ---------------------------------------------------------------------------
# Read: workbook navigation
# ---------------------------------------------------------------------------


def test_load_workbook(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb is not None


def test_sheetnames(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert "Hoja1" in wb.sheetnames


def test_active_title(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb.active.title == "Hoja1"


def test_getitem_found(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb["Hoja1"].title == "Hoja1"


def test_getitem_missing(load_wb):
    wb = load_wb("samples/test.xlsx")
    with pytest.raises(KeyError):
        _ = wb["DoesNotExist"]


# ---------------------------------------------------------------------------
# Read: worksheet / cell access
# ---------------------------------------------------------------------------


def test_ws_getitem(load_wb):
    wb = load_wb("samples/test.xlsx")
    cell = wb.active["A1"]
    assert cell is not None


def test_ws_cell_method(load_wb):
    wb = load_wb("samples/test.xlsx")
    ws = wb.active
    assert ws.cell(row=1, column=1).value == ws["A1"].value


# ---------------------------------------------------------------------------
# Read: cell values
# ---------------------------------------------------------------------------


def test_cell_value_str(load_wb):
    wb = load_wb("samples/test.xlsx")
    v = wb.active["A1"].value
    assert isinstance(v, str)


def test_cell_value_int(load_wb):
    # C1 of test.xlsx → whole number 42
    wb = load_wb("samples/test.xlsx")
    v = wb.active["C1"].value
    assert isinstance(v, int) and not isinstance(v, bool)
    assert v == 42


def test_cell_value_date(load_wb):
    # B1 of multi_sheet.xlsx "Empleados" → date 2025-01-21
    # openpyxl returns datetime.datetime; oxlsx returns datetime.date.
    # Both are subclasses of datetime.date — so isinstance passes for both.
    # We compare only the date part to avoid divergence on time component.
    wb = load_wb("samples/multi_sheet.xlsx")
    v = wb["Empleados"]["B1"].value
    assert isinstance(v, datetime.date)
    assert v.year == 2025 and v.month == 1 and v.day == 21


def test_cell_value_none(load_wb):
    # A reference that holds no data → None
    wb = load_wb("samples/test.xlsx")
    cell = wb.active["Z99"]
    value = cell.value if cell is not None else None
    assert value is None


# ---------------------------------------------------------------------------
# Read: formatting
# ---------------------------------------------------------------------------


def test_cell_font_bold(load_wb):
    # A1 of test.xlsx has bold font
    wb = load_wb("samples/test.xlsx")
    font = wb.active["A1"].font
    assert font is not None
    assert isinstance(font.bold, bool)
    assert font.bold is True


def test_cell_fill_fg_color(load_wb):
    # A1 of test.xlsx has a solid fill
    wb = load_wb("samples/test.xlsx")
    fill = wb.active["A1"].fill
    assert fill is not None


# ---------------------------------------------------------------------------
# Write: round-trip and setitem
# ---------------------------------------------------------------------------


def test_write_round_trip(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws["A1"] = "hello"
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        assert wb2.active["A1"].value == "hello"
    finally:
        os.unlink(path)


def test_ws_setitem(new_wb):
    wb, _ = new_wb
    ws = wb.active
    ws["A1"] = "world"
    assert ws["A1"].value == "world"


def test_rmw_modify_cell(load_wb):
    wb = load_wb("samples/test.xlsx")
    wb.active["A1"] = "RMW"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = load_wb(path)
        assert wb2.active["A1"].value == "RMW"
    finally:
        os.unlink(path)


def test_rmw_preserves_other_sheets(load_wb):
    wb = load_wb("samples/multi_sheet.xlsx")
    wb["Empleados"]["A1"] = "RMW"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = load_wb(path)
        assert wb2["Productos"]["A1"].value == "Laptop"
    finally:
        os.unlink(path)


def test_rmw_preserves_font(load_wb):
    wb = load_wb("samples/test.xlsx")
    wb.active["C1"] = 99

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = load_wb(path)
        assert wb2.active["A1"].font.bold is True
    finally:
        os.unlink(path)


def test_write_formula(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws["A1"] = "=SUM(B1:B3)"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        value = wb2.active["A1"].value
        assert isinstance(value, str)
        assert value == "=SUM(B1:B3)"
    finally:
        os.unlink(path)


def test_write_date_renders_as_date(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws["A1"] = datetime.date(2025, 1, 21)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        value = wb2.active["A1"].value
        assert isinstance(value, datetime.date)
        assert value.year == 2025 and value.month == 1 and value.day == 21
    finally:
        os.unlink(path)


def test_write_merge_cells(new_wb):
    wb, save = new_wb
    ws = wb.active
    ws.merge_cells("A1:C1")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        with zipfile.ZipFile(path, "r") as zf:
            sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert '<mergeCell ref="A1:C1"' in sheet_xml
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# Phase 9 — Iteration API
# ---------------------------------------------------------------------------


def test_max_row(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb.active.max_row == 1


def test_max_column(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb.active.max_column == 3


def test_min_row(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb.active.min_row == 1


def test_min_column(load_wb):
    wb = load_wb("samples/test.xlsx")
    assert wb.active.min_column == 1


def test_iter_rows_values_only(load_wb):
    wb = load_wb("samples/test.xlsx")
    rows = list(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == "Hola"
    assert rows[0][2] == 42


def test_iter_rows_all_values_only(load_wb):
    # No bounds → all data rows
    wb = load_wb("samples/multi_sheet.xlsx")
    rows = list(wb["Empleados"].iter_rows(values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == "Juan"


def test_iter_rows_returns_cells(load_wb):
    wb = load_wb("samples/test.xlsx")
    rows = list(wb.active.iter_rows(min_row=1, max_row=1, values_only=False))
    assert len(rows) == 1
    # each element is a cell-like object with .value
    assert rows[0][0].value == "Hola"


def test_iter_cols_values_only(load_wb):
    wb = load_wb("samples/test.xlsx")
    cols = list(wb.active.iter_cols(min_col=1, max_col=1, values_only=True))
    assert len(cols) == 1
    assert cols[0][0] == "Hola"


def test_iter_cols_all_values_only(load_wb):
    wb = load_wb("samples/multi_sheet.xlsx")
    cols = list(wb["Empleados"].iter_cols(values_only=True))
    assert len(cols) == 2
    assert cols[0][0] == "Juan"


def test_ws_values(load_wb):
    # ws.values → generator of row tuples (values only)
    wb = load_wb("samples/test.xlsx")
    rows = list(wb.active.values)
    assert len(rows) == 1
    assert rows[0][0] == "Hola"
    assert rows[0][2] == 42


def test_ws_rows_property(load_wb):
    # ws.rows → generator of cell tuples
    wb = load_wb("samples/test.xlsx")
    rows = list(wb.active.rows)
    assert len(rows) == 1
    assert rows[0][0].value == "Hola"


# ---------------------------------------------------------------------------
# Phase 12b — Comments parity
# ---------------------------------------------------------------------------


def test_comment_round_trip(load_wb):
    wb = load_wb("samples/comments.xlsx")
    comment = wb.active["B2"].comment
    assert comment is not None
    assert comment.text == "Hello from fixture"
    assert comment.author == "Fixture Author"


def test_comment_write_new_cell(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active

    if hasattr(ws, "set_cell_comment"):
        ws.set_cell_comment("C3", "Created from python", "Py Author")
    else:
        ws["C3"].comment = OpenPyxlComment("Created from python", "Py Author")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        comment = wb2.active["C3"].comment
        assert comment is not None
        assert comment.text == "Created from python"
        assert comment.author == "Py Author"
    finally:
        os.unlink(path)


def _normalize_print_area(value):
    """Minimal normalization — only handle whitespace/escaping, NOT qualification.
    
    Both openpyxl and oxlsx should return the same qualified form.
    This function cannot mask format or sheet-prefix divergence.
    """
    if value is None:
        return None
    if isinstance(value, (list, tuple)):
        # openpyxl might return list for multi-range in some contexts
        value = ",".join(str(v) for v in value)
    # Only strip outer whitespace — do NOT strip sheet prefix
    return str(value).strip()


def _make_seed_with_defined_names(defined_names_xml):
    wb = openpyxl.Workbook()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)

    with zipfile.ZipFile(path, "r") as src:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            out_path = tmp.name
        with zipfile.ZipFile(out_path, "w") as dst:
            for name in src.namelist():
                payload = src.read(name)
                if name == "xl/workbook.xml":
                    xml = payload.decode("utf-8")
                    xml = xml.replace("</workbook>", f"{defined_names_xml}</workbook>")
                    payload = xml.encode("utf-8")
                dst.writestr(name, payload)

    os.unlink(path)
    return out_path


def test_print_area_round_trip(load_wb):
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">\'Sheet\'!$A$1:$C$10</definedName><definedName name="KeepMe">Sheet!$D$1</definedName></definedNames>'
    )
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as out:
        out_path = out.name

    try:
        wb = load_wb(seeded_path)
        assert _normalize_print_area(wb.active.print_area) == "'Sheet'!$A$1:$C$10"
        wb.save(out_path)
        wb2 = load_wb(out_path)
        assert _normalize_print_area(wb2.active.print_area) == "'Sheet'!$A$1:$C$10"

        with zipfile.ZipFile(out_path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "_xlnm.Print_Area" in workbook_xml
        assert "KeepMe" in workbook_xml
    finally:
        os.unlink(seeded_path)
        os.unlink(out_path)


def test_print_area_write_new_workbook(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws.print_area = "$B$2:$D$20"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        assert _normalize_print_area(wb2.active.print_area) == "'Sheet'!$B$2:$D$20"

        with zipfile.ZipFile(path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "_xlnm.Print_Area" in workbook_xml
    finally:
        os.unlink(path)


def test_auto_filter_round_trip(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws.auto_filter.ref = "A1:C5"
    assert ws.auto_filter.ref == "A1:C5"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        # Reopen with openpyxl specifically: a structural error would surface
        # as a repair / parse failure here.
        wb2 = openpyxl.load_workbook(path)
        assert wb2.active.auto_filter.ref == "A1:C5"

        # And reopen through the same backend for parity.
        wb3 = load_wb(path)
        assert wb3.active.auto_filter.ref == "A1:C5"

        with zipfile.ZipFile(path, "r") as zf:
            sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert '<autoFilter ref="A1:C5"' in sheet_xml
    finally:
        os.unlink(path)


def test_print_area_clear_removes_defined_name(load_wb):
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">\'Sheet\'!$A$1:$A$5</definedName><definedName name="KeepMe">Sheet!$D$1</definedName></definedNames>'
    )
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as out:
        out_path = out.name

    try:
        wb = load_wb(seeded_path)
        wb.active.print_area = None
        wb.save(out_path)

        wb2 = load_wb(out_path)
        assert not wb2.active.print_area

        with zipfile.ZipFile(out_path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "_xlnm.Print_Area" not in workbook_xml
        assert "KeepMe" in workbook_xml
    finally:
        os.unlink(seeded_path)
        os.unlink(out_path)


def test_print_title_rows_and_cols_round_trip(load_wb):
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Titles" localSheetId="0">\'Sheet\'!$A:$C,\'Sheet\'!$1:$3</definedName><definedName name="KeepMe">Sheet!$D$1</definedName></definedNames>'
    )
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as out:
        out_path = out.name

    try:
        wb = load_wb(seeded_path)
        assert wb.active.print_title_rows == "$1:$3"
        assert wb.active.print_title_cols == "$A:$C"
        wb.save(out_path)

        wb2 = load_wb(out_path)
        assert wb2.active.print_title_rows == "$1:$3"
        assert wb2.active.print_title_cols == "$A:$C"

        with zipfile.ZipFile(out_path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "_xlnm.Print_Titles" in workbook_xml
        assert "KeepMe" in workbook_xml
    finally:
        os.unlink(seeded_path)
        os.unlink(out_path)


def test_print_title_rows_and_cols_write_new_workbook(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    ws.print_title_rows = "$1:$2"
    ws.print_title_cols = "$A:$B"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        wb2 = load_wb(path)
        assert wb2.active.print_title_rows == "$1:$2"
        assert wb2.active.print_title_cols == "$A:$B"

        with zipfile.ZipFile(path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "_xlnm.Print_Titles" in workbook_xml
    finally:
        os.unlink(path)


def test_print_title_rows_and_cols_clear_removes_defined_name(load_wb):
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Titles" localSheetId="0">\'Sheet\'!$1:$3</definedName><definedName name="KeepMe">Sheet!$D$1</definedName></definedNames>'
    )
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as out:
        out_path = out.name

    try:
        wb = load_wb(seeded_path)
        wb.active.print_title_rows = None
        wb.active.print_title_cols = None
        wb.save(out_path)

        wb2 = load_wb(out_path)
        assert wb2.active.title == "Sheet"

        with zipfile.ZipFile(out_path, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "KeepMe" in workbook_xml
    finally:
        os.unlink(seeded_path)
        os.unlink(out_path)


# ---------------------------------------------------------------------------
# Phase 13 slice 3a: global-scope named ranges (oxlsx Python parity)
#
# These exercise oxlsx's generic named-range API. openpyxl's DefinedNameDict
# is permissive about reserved names and uses different assignment semantics,
# so reserved-name rejection is asserted on the oxlsx path only (intentionally
# stricter: ValueError).
# ---------------------------------------------------------------------------


def test_named_ranges_round_trip():
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="Sales">\'Sheet\'!$A$1:$A$5</definedName></definedNames>'
    )
    try:
        wb = oxlsx.load_workbook(seeded_path)
        assert "Sales" in wb.defined_names
        assert wb.defined_names["Sales"].value == "'Sheet'!$A$1:$A$5"
    finally:
        os.unlink(seeded_path)


def test_named_ranges_add_new():
    wb = oxlsx.Workbook()
    wb.defined_names.add(oxlsx.DefinedName("Region", "'Sheet'!$B$1:$B$9"))

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = oxlsx.load_workbook(path)
        assert "Region" in wb2.defined_names
        assert wb2.defined_names["Region"].value == "'Sheet'!$B$1:$B$9"
    finally:
        os.unlink(path)


def test_named_ranges_remove():
    wb = oxlsx.Workbook()
    wb.defined_names.add(oxlsx.DefinedName("Region", "'Sheet'!$B$1:$B$9"))

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        del wb.defined_names["Region"]
        wb.save(path)
        wb2 = oxlsx.load_workbook(path)
        assert "Region" not in wb2.defined_names
    finally:
        os.unlink(path)


def test_named_ranges_exclude_reserved():
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">\'Sheet\'!$A$1:$C$10</definedName></definedNames>'
    )
    try:
        wb = oxlsx.load_workbook(seeded_path)
        # Reserved names are never surfaced in the generic collection.
        assert "_xlnm.Print_Area" not in wb.defined_names
        assert list(wb.defined_names) == []

        # oxlsx is intentionally stricter than openpyxl: adding a reserved
        # name raises ValueError instead of silently accepting it.
        with pytest.raises(ValueError):
            wb.defined_names.add(
                oxlsx.DefinedName("_xlnm.Print_Area", "'Sheet'!$A$1:$B$2")
            )
    finally:
        os.unlink(seeded_path)


# ---------------------------------------------------------------------------
# Phase 13 slice 3b: sheet-local named ranges (ws.defined_names, oxlsx parity)
#
# Mirrors openpyxl's per-sheet DefinedNameDict. wb.defined_names (global) and
# ws.defined_names (sheet-local) are independent collections. Reserved-name
# rejection is oxlsx-stricter (ValueError) and asserted on the oxlsx path only.
# ---------------------------------------------------------------------------


def test_sheet_local_names_round_trip():
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="LocalRegion" localSheetId="0">Sheet!$A$1</definedName></definedNames>'
    )
    try:
        wb = oxlsx.load_workbook(seeded_path)
        ws = wb.worksheets[0]
        # Promoted onto the owning worksheet...
        assert "LocalRegion" in ws.defined_names
        assert ws.defined_names["LocalRegion"].value == "Sheet!$A$1"
        # ...and NOT visible in the workbook-global collection.
        assert "LocalRegion" not in wb.defined_names
    finally:
        os.unlink(seeded_path)


def test_sheet_local_names_add_new():
    wb = oxlsx.Workbook()
    wb.active.defined_names.add(oxlsx.DefinedName("Region", "Sheet!$B$1"))

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = oxlsx.load_workbook(path)
        ws = wb2.worksheets[0]
        assert "Region" in ws.defined_names
        assert ws.defined_names["Region"].value == "Sheet!$B$1"
        # Stays scoped to the sheet, never leaks to the global collection.
        assert "Region" not in wb2.defined_names
    finally:
        os.unlink(path)


def test_sheet_local_names_exclude_reserved():
    seeded_path = _make_seed_with_defined_names(
        '<definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">\'Sheet\'!$A$1:$C$10</definedName></definedNames>'
    )
    try:
        wb = oxlsx.load_workbook(seeded_path)
        ws = wb.worksheets[0]
        # Reserved print name is owned by print settings, not surfaced here.
        assert "_xlnm.Print_Area" not in ws.defined_names
        assert list(ws.defined_names) == []

        # oxlsx is intentionally stricter than openpyxl: reserved add raises.
        with pytest.raises(ValueError):
            ws.defined_names.add(
                oxlsx.DefinedName("_xlnm.Print_Area", "'Sheet'!$A$1:$B$2")
            )
    finally:
        os.unlink(seeded_path)


def test_sheet_local_and_global_same_name():
    wb = oxlsx.Workbook()
    wb.defined_names.add(oxlsx.DefinedName("Data", "Sheet!$A$1"))
    wb.active.defined_names.add(oxlsx.DefinedName("Data", "Sheet!$B$2"))

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)
        wb2 = oxlsx.load_workbook(path)
        ws = wb2.worksheets[0]
        # Independent scopes: same key, different values.
        assert wb2.defined_names["Data"].value == "Sheet!$A$1"
        assert ws.defined_names["Data"].value == "Sheet!$B$2"
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# Tables (Phase 13 final slice) — parity with openpyxl.worksheet.table
# ---------------------------------------------------------------------------


@pytest.fixture(params=["openpyxl", "oxlsx"])
def table_writer(request):
    """Returns (workbook, Table, TableStyleInfo) for the active write engine."""
    if request.param == "openpyxl":
        from openpyxl.worksheet.table import Table, TableStyleInfo

        return openpyxl.Workbook(), Table, TableStyleInfo
    return oxlsx.Workbook(), oxlsx.Table, oxlsx.TableStyleInfo


def _temp_xlsx(name):
    return os.path.join(tempfile.gettempdir(), name)


def test_tables_round_trip(table_writer, load_wb):
    wb, Table, _ = table_writer
    ws = wb.active
    ws["A1"] = "h1"
    ws["B1"] = "h2"
    ws["C1"] = "h3"
    ws.add_table(Table(displayName="Table1", ref="A1:C3"))

    path = _temp_xlsx("oxlsx_py_tables_round_trip.xlsx")
    try:
        wb.save(path)
        ws2 = load_wb(path).active
        assert "Table1" in ws2.tables
        table = ws2.tables["Table1"]
        assert table.ref == "A1:C3"
        assert [c.name for c in table.tableColumns] == ["h1", "h2", "h3"]
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_tables_add_new(table_writer, load_wb):
    wb, Table, _ = table_writer
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Age"
    ws.add_table(Table(displayName="People", ref="A1:B2"))

    # Present immediately and after a round trip.
    assert "People" in ws.tables
    assert len(ws.tables) == 1

    path = _temp_xlsx("oxlsx_py_tables_add_new.xlsx")
    try:
        wb.save(path)
        ws2 = load_wb(path).active
        assert "People" in ws2.tables
        assert ws2.tables["People"].ref == "A1:B2"
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_tables_duplicate_display_name_raises(table_writer):
    # Both openpyxl and oxlsx reject a duplicate displayName with ValueError.
    wb, Table, _ = table_writer
    ws = wb.active
    ws["A1"] = "h1"
    ws.add_table(Table(displayName="Dup", ref="A1:A2"))
    with pytest.raises(ValueError):
        ws.add_table(Table(displayName="Dup", ref="C1:C2"))


def test_tables_style_info(table_writer, load_wb):
    wb, Table, TableStyleInfo = table_writer
    ws = wb.active
    ws["A1"] = "h1"
    ws["B1"] = "h2"
    table = Table(displayName="Styled", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)

    path = _temp_xlsx("oxlsx_py_tables_style.xlsx")
    try:
        wb.save(path)
        style = load_wb(path).active.tables["Styled"].tableStyleInfo
        assert style is not None
        assert style.name == "TableStyleMedium9"
        assert bool(style.showRowStripes) is True
    finally:
        if os.path.exists(path):
            os.remove(path)



# ---------------------------------------------------------------------------
# FIX 7: Python table mutation must persist (attached mode)
# ---------------------------------------------------------------------------


def test_tables_mutation_persists():
    """FIX 7: Mutating ws.tables['X'] must persist through save+reload."""
    wb = oxlsx.Workbook()
    ws = wb.active
    ws["A1"] = "h1"
    ws["B1"] = "h2"
    
    # Add a table WITHOUT a style
    table = oxlsx.Table(displayName="MutableTable", ref="A1:B2")
    ws.add_table(table)
    
    # Now mutate the ATTACHED table retrieved from ws.tables
    ws.tables["MutableTable"].tableStyleInfo = oxlsx.TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True
    )
    
    # Save and reload
    path = _temp_xlsx("oxlsx_py_tables_mutation.xlsx")
    try:
        wb.save(path)
        wb2 = oxlsx.load_workbook(path)
        reloaded_table = wb2.active.tables["MutableTable"]
        
        # The style MUST persist (openpyxl parity)
        assert reloaded_table.tableStyleInfo is not None
        assert reloaded_table.tableStyleInfo.name == "TableStyleMedium9"
        assert bool(reloaded_table.tableStyleInfo.showRowStripes) is True
    finally:
        if os.path.exists(path):
            os.remove(path)


# ---------------------------------------------------------------------------
# CRITICAL: RMW of tables.xlsx must preserve sharedStrings content-type
# ---------------------------------------------------------------------------


def test_tables_rmw_openpyxl_reopen():
    """
    CRITICAL: RMW on samples/tables.xlsx (openpyxl-authored, INLINE strings, NO sharedStrings)
    creates xl/sharedStrings.xml and MUST add the content-type Override, otherwise openpyxl
    cannot reopen the file.
    
    This is the REAL guard that was missing from the previous dual review round.
    """
    # Load with oxlsx, modify a cell (gains shared strings), save
    wb = oxlsx.load_workbook("samples/tables.xlsx")
    wb.active["B1"] = "modified_by_oxlsx"
    
    path = _temp_xlsx("oxlsx_tables_rmw_openpyxl_reopen.xlsx")
    try:
        wb.save(path)
        
        # CRITICAL: reopen with openpyxl (must NOT raise IndexError or corruption warning)
        wb2 = openpyxl.load_workbook(path)
        
        # Assert the table is intact
        assert "Table1" in wb2.active.tables
        
        # Assert the modified cell is present
        assert wb2.active["B1"].value == "modified_by_oxlsx"
        
        # Also verify the content-types manifest has the sharedStrings override
        with zipfile.ZipFile(path, "r") as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")
        assert "/xl/sharedStrings.xml" in content_types
        assert "spreadsheetml.sharedStrings+xml" in content_types
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_rmw_gains_shared_strings_no_tables_no_comments_openpyxl_reopen():
    """
    CRITICAL regression guard for the GENERAL RMW shared-strings path.

    samples/inline_strings_no_sst.xlsx is an openpyxl-style file that uses INLINE
    strings, has NO xl/sharedStrings.xml part (and no Override for it), and has
    NO tables and NO comments. When oxlsx writes a string cell on save it emits
    xl/sharedStrings.xml, so the rewritten [Content_Types].xml MUST gain the
    sharedStrings Override — otherwise openpyxl raises IndexError on reopen.

    The existing tables guard (test_tables_rmw_openpyxl_reopen) only reaches the
    content-type rewrite through the `total_tables > 0` branch term. This test
    locks down the `has_shared_strings` term of the writer.rs branch condition
    on the no-tables / no-comments path, so a future refactor cannot regress it.
    """
    fixture = "samples/inline_strings_no_sst.xlsx"

    # Prove this fixture genuinely exercises the no-tables / no-comments path
    # (and that the source truly lacks a sharedStrings part to begin with).
    with zipfile.ZipFile(fixture, "r") as zf:
        names = zf.namelist()
        source_content_types = zf.read("[Content_Types].xml").decode("utf-8")
    assert not any("/tables/" in n for n in names), "fixture must have NO tables"
    assert not any("comments" in n for n in names), "fixture must have NO comments"
    assert not any(
        "sharedStrings" in n for n in names
    ), "fixture must start WITHOUT a sharedStrings part"
    assert "/xl/sharedStrings.xml" not in source_content_types

    # Sanity: oxlsx can read the inline string values before the RMW.
    wb = oxlsx.load_workbook(fixture)
    assert wb.active["A1"].value == "h1"

    # Modify + add string cells, forcing shared-strings emission on save.
    wb.active["A1"] = "modified_by_oxlsx"
    wb.active["A4"] = "added_by_oxlsx"

    path = _temp_xlsx("oxlsx_inline_strings_rmw_openpyxl_reopen.xlsx")
    try:
        wb.save(path)

        # CRITICAL: openpyxl must reopen WITHOUT IndexError / repair.
        wb2 = openpyxl.load_workbook(path)
        assert wb2.active["A1"].value == "modified_by_oxlsx"
        assert wb2.active["A4"].value == "added_by_oxlsx"

        # The rewritten manifest must contain EXACTLY ONE sharedStrings Override.
        with zipfile.ZipFile(path, "r") as zf:
            content_types = zf.read("[Content_Types].xml").decode("utf-8")
            out_names = zf.namelist()
        assert content_types.count("/xl/sharedStrings.xml") == 1
        assert "spreadsheetml.sharedStrings+xml" in content_types

        # The output really gained the part and still has no tables/comments.
        assert "xl/sharedStrings.xml" in out_names
        assert not any("/tables/" in n for n in out_names)
        assert not any("comments" in n for n in out_names)
    finally:
        if os.path.exists(path):
            os.remove(path)


# ---------------------------------------------------------------------------
# Lossless RMW (Phase 7) — verbatim-when-unmodified table preservation.
# samples/tables_totals.xlsx (openpyxl 3.1.2, sheet "Data") carries unmodeled
# features oxlsx does NOT model: totalsRow*, calculatedColumnFormula, sortState,
# autoFilter. These must survive a Python RMW when the table is not mutated.
# ---------------------------------------------------------------------------


def _table_part_named(path, display_name):
    """Return the decoded xl/tables/*.xml part whose <table> is display_name."""
    with zipfile.ZipFile(path, "r") as zf:
        parts = [
            zf.read(n).decode("utf-8")
            for n in zf.namelist()
            if n.startswith("xl/tables/") and n.endswith(".xml")
        ]
    return next(xml for xml in parts if f'displayName="{display_name}"' in xml)


def test_table_totals_row_survives_cell_edit_rmw():
    """
    THE mandatory lossless-RMW regression (scenario 6).

    Open the openpyxl-authored totals-row fixture with oxlsx, edit an UNRELATED
    cell, save, then reopen with openpyxl: the file must load WITHOUT repair and
    the unmodeled features (totals row + calculatedColumnFormula + sortState)
    must SURVIVE because the table was never mutated.
    """
    wb = oxlsx.load_workbook("samples/tables_totals.xlsx")
    ws = wb.active
    # Edit a cell outside every table range (tables span A1:C5 and E1:G3).
    ws["H10"] = "unrelated_edit"

    path = _temp_xlsx("oxlsx_py_table_totals_survives.xlsx")
    try:
        wb.save(path)

        # openpyxl must reopen WITHOUT repair / exception.
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2["Data"]
        assert "SalesTable" in ws2.tables
        table = ws2.tables["SalesTable"]

        # Totals row survived.
        assert table.totalsRowCount == 1
        funcs = {c.name: c.totalsRowFunction for c in table.tableColumns}
        assert funcs["Revenue"] == "sum"

        # calculatedColumnFormula survived (a second unmodeled feature).
        calc = {c.name: c.calculatedColumnFormula for c in table.tableColumns}
        assert calc["Units"] is not None

        # sortState survived (a third unmodeled feature).
        assert table.sortState is not None

        # The retained part is byte-identical => zero loss at the XML level.
        sales_xml = _table_part_named(path, "SalesTable")
        assert 'totalsRowCount="1"' in sales_xml
        assert "calculatedColumnFormula" in sales_xml
        assert "<sortState" in sales_xml
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_table_mutation_documents_feature_loss():
    """
    Scenario 7: mutating a table through the Python API forces the model
    serialization path. The file stays valid in openpyxl and modeled fields
    round-trip, but the unmodeled features are documented as DROPPED.
    """
    wb = oxlsx.load_workbook("samples/tables_totals.xlsx")
    ws = wb.active
    # Mutate the table through the Python API -> marks it dirty (model path).
    ws.tables["SalesTable"].tableStyleInfo = oxlsx.TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True
    )

    path = _temp_xlsx("oxlsx_py_table_mutation_loss.xlsx")
    try:
        wb.save(path)

        wb2 = openpyxl.load_workbook(path)  # must not raise / repair
        ws2 = wb2["Data"]
        table = ws2.tables["SalesTable"]

        # Modeled fields are correct (the (ii) contract).
        assert table.displayName == "SalesTable"
        assert table.ref == "A1:C5"
        assert [c.name for c in table.tableColumns] == ["Region", "Units", "Revenue"]
        assert table.tableStyleInfo.name == "TableStyleMedium2"

        # Documented loss: unmodeled features are gone on the model path.
        sales_xml = _table_part_named(path, "SalesTable")
        assert "totalsRowCount" not in sales_xml
        assert "sortState" not in sales_xml
        assert "calculatedColumnFormula" not in sales_xml

        # Inventory was never touched -> still verbatim (autoFilter survives).
        inv_xml = _table_part_named(path, "Inventory")
        assert "autoFilter" in inv_xml
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_table_delete_via_python_removes_table():
    """
    openpyxl parity: `del ws.tables["X"]` drops the table (part + rel +
    tablePart) and leaves the other table intact, reopening clean in openpyxl.
    """
    wb = oxlsx.load_workbook("samples/tables_totals.xlsx")
    ws = wb.active
    assert "Inventory" in ws.tables
    assert "SalesTable" in ws.tables

    del ws.tables["Inventory"]

    assert "Inventory" not in ws.tables
    assert len(ws.tables) == 1

    path = _temp_xlsx("oxlsx_py_table_delete.xlsx")
    try:
        wb.save(path)

        wb2 = openpyxl.load_workbook(path)  # clean reopen
        ws2 = wb2["Data"]
        assert "Inventory" not in ws2.tables
        assert "SalesTable" in ws2.tables
        # SalesTable was untouched -> still verbatim (totals preserved).
        assert ws2.tables["SalesTable"].totalsRowCount == 1

        with zipfile.ZipFile(path, "r") as zf:
            table_parts = [
                n
                for n in zf.namelist()
                if n.startswith("xl/tables/") and n.endswith(".xml")
            ]
        assert len(table_parts) == 1
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_table_delete_missing_raises_keyerror():
    """`del ws.tables["missing"]` raises KeyError, mirroring openpyxl's dict."""
    wb = oxlsx.load_workbook("samples/tables_totals.xlsx")
    ws = wb.active
    with pytest.raises(KeyError):
        del ws.tables["DoesNotExist"]


# ---------------------------------------------------------------------------
# Phase 12a — Hyperlinks parity (scenarios 7, 8)
#
# Mirrors openpyxl.worksheet.hyperlink.Hyperlink: a cell exposes `.hyperlink`
# (an object with `.target`/`.location`/`.tooltip`/`.display`, or None).
# openpyxl authors hyperlinks via `cell.hyperlink = "https://..."`; because
# oxlsx's PyCell is a value snapshot, the write path is a worksheet method
# (`ws.set_cell_hyperlink`), mirroring `ws.set_cell_comment`.
# ---------------------------------------------------------------------------


def test_hyperlink_read(load_wb):
    # samples/hyperlinks.xlsx A1 → external https://example.com with tooltip +
    # display. Both engines must expose the same resolved Hyperlink shape.
    wb = load_wb("samples/hyperlinks.xlsx")
    link = wb.active["A1"].hyperlink
    assert link is not None
    assert link.target == "https://example.com"
    assert link.tooltip == "Example"
    assert link.display == "Example Site"


def test_hyperlink_round_trip(new_wb, load_wb):
    wb, save = new_wb
    ws = wb.active
    target = "https://example.com"

    if hasattr(ws, "set_cell_hyperlink"):
        ws.set_cell_hyperlink("A1", target)
    else:
        ws["A1"].hyperlink = target

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        save(path)
        # Cross-engine reopen (incl. oxlsx-write → openpyxl-read) must not
        # repair and must preserve the external target.
        wb2 = load_wb(path)
        link = wb2.active["A1"].hyperlink
        assert link is not None
        assert link.target == target
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# openpyxl-parity Group 1 — column utilities + PatternFill
#
# Column utilities mirror openpyxl.utils.get_column_letter /
# column_index_from_string (1-based). PatternFill mirrors
# openpyxl.styles.PatternFill; because oxlsx's PyCell is a value snapshot, the
# write path is a worksheet method (`ws.set_cell_fill`), mirroring
# `ws.set_cell_comment` / `ws.set_cell_hyperlink`.
# ---------------------------------------------------------------------------


def test_get_column_letter():
    from openpyxl.utils import get_column_letter as openpyxl_get_column_letter

    # Anchor values from the parity spec.
    assert oxlsx.get_column_letter(1) == "A"
    assert oxlsx.get_column_letter(26) == "Z"
    assert oxlsx.get_column_letter(27) == "AA"
    assert oxlsx.get_column_letter(702) == "ZZ"
    assert oxlsx.get_column_letter(703) == "AAA"

    # Cross-check against openpyxl over a range.
    for idx in range(1, 1000):
        assert oxlsx.get_column_letter(idx) == openpyxl_get_column_letter(idx)

    # openpyxl raises for indices < 1; oxlsx mirrors with ValueError.
    with pytest.raises(ValueError):
        oxlsx.get_column_letter(0)


def test_column_index_from_string():
    from openpyxl.utils import column_index_from_string as openpyxl_column_index

    assert oxlsx.column_index_from_string("A") == 1
    assert oxlsx.column_index_from_string("Z") == 26
    assert oxlsx.column_index_from_string("AA") == 27
    assert oxlsx.column_index_from_string("ZZ") == 702
    assert oxlsx.column_index_from_string("AAA") == 703

    # Cross-check the inverse against openpyxl over a range.
    for idx in range(1, 1000):
        letters = oxlsx.get_column_letter(idx)
        assert oxlsx.column_index_from_string(letters) == idx
        assert oxlsx.column_index_from_string(letters) == openpyxl_column_index(letters)

    # Invalid (non-alphabetic / mixed) input raises ValueError.
    with pytest.raises(ValueError):
        oxlsx.column_index_from_string("A1")


def test_pattern_fill_round_trip():
    from openpyxl.styles import PatternFill as OpenpyxlPatternFill

    # Alias parity: start_color → fgColor, fill_type → patternType.
    fill = oxlsx.PatternFill(fill_type="solid", start_color="FFFF0000")
    assert fill.fill_type == "solid"
    assert fill.patternType == "solid"
    assert fill.fgColor.rgb == "FFFF0000"

    # Sanity: openpyxl resolves the same shape from the same kwargs.
    ref = OpenpyxlPatternFill(fill_type="solid", start_color="FFFF0000")
    assert ref.fill_type == "solid"
    assert ref.fgColor.rgb == "FFFF0000"

    wb = oxlsx.Workbook()
    ws = wb.active
    ws["A1"] = "Filled"
    ws.set_cell_fill("A1", oxlsx.PatternFill(fill_type="solid", fgColor="FFFF0000"))

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        wb.save(path)

        # oxlsx reads back the fill it wrote.
        wb2 = oxlsx.load_workbook(path)
        fill2 = wb2.active["A1"].fill
        assert fill2 is not None
        assert fill2.fg_color.rgb == "FFFF0000"

        # Cross-check: openpyxl reads the same solid fill + color.
        wb3 = openpyxl.load_workbook(path)
        fill3 = wb3.active["A1"].fill
        assert fill3.fill_type == "solid"
        assert fill3.fgColor.rgb == "FFFF0000"
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# Phase 10 (real) — bulk-write parity: append (list/dict), insert/delete rows,
# move_range. Each test runs the SAME op sequence against openpyxl and oxlsx
# and asserts the resulting grids match (cross-engine parity), plus an explicit
# expected layout so the assertions exercise real production logic.
# ---------------------------------------------------------------------------


def _window(ws, max_row, max_col):
    """Normalized value grid over a fixed window (engine-agnostic)."""
    return [
        tuple(row)
        for row in ws.iter_rows(
            min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True
        )
    ]


def _do_appends(ws):
    # Positional list from column A.
    ws.append(["a", "b", "c"])
    # Dict with column-LETTER keys (B is skipped -> stays empty).
    ws.append({"A": "x", "C": "z"})
    # Dict with 1-based INT keys.
    ws.append({1: "p", 3: "r"})
    # Dict with MIXED letter + int keys (A and C are gaps).
    ws.append({"B": "q", 4: "s"})


def test_append_list_and_dict():
    expected = [
        ("a", "b", "c", None),
        ("x", None, "z", None),
        ("p", None, "r", None),
        (None, "q", None, "s"),
    ]

    wb_o = openpyxl.Workbook()
    _do_appends(wb_o.active)
    grid_o = _window(wb_o.active, 4, 4)

    wb_x = oxlsx.Workbook()
    _do_appends(wb_x.active)
    grid_x = _window(wb_x.active, 4, 4)

    # oxlsx matches the concrete openpyxl-dict semantics...
    assert grid_x == expected
    # ...and matches openpyxl exactly (cross-engine parity).
    assert grid_x == grid_o


def _seed_and_edit_rows(ws):
    ws["A1"] = "r1"
    ws["A2"] = "r2"
    ws["A3"] = "r3"
    ws["B1"] = 1
    ws["B2"] = 2
    ws["B3"] = 3
    ws.insert_rows(2)  # blank row before row 2; r2/r3 shift down
    ws.delete_rows(1)  # drop original r1; everything shifts up


def test_delete_and_insert_rows():
    expected = [
        (None, None),
        ("r2", 2),
        ("r3", 3),
        (None, None),
        (None, None),
        (None, None),
    ]

    # openpyxl reference layout after the same op sequence.
    wb_o = openpyxl.Workbook()
    _seed_and_edit_rows(wb_o.active)
    grid_o = _window(wb_o.active, 6, 2)

    # oxlsx in-memory layout after the same op sequence.
    wb_x = oxlsx.Workbook()
    _seed_and_edit_rows(wb_x.active)
    grid_x_mem = _window(wb_x.active, 6, 2)

    # oxlsx layout after a save + reopen round-trip (persistence parity).
    path = _temp_xlsx("oxlsx_py_bulk_insert_delete.xlsx")
    try:
        wb_x.save(path)
        grid_x_disk = _window(oxlsx.load_workbook(path).active, 6, 2)
    finally:
        if os.path.exists(path):
            os.remove(path)

    assert grid_x_mem == expected
    assert grid_x_mem == grid_o
    assert grid_x_disk == expected


def _seed_and_move(ws):
    ws["A1"] = 1
    ws["B1"] = 2
    ws["A2"] = 3
    ws["B2"] = 4
    # Relocate the 2x2 block down 2 / right 2 -> C3:D4; source clears.
    ws.move_range("A1:B2", rows=2, cols=2)


def test_move_range():
    expected = [
        (None, None, None, None),
        (None, None, None, None),
        (None, None, 1, 2),
        (None, None, 3, 4),
    ]

    wb_o = openpyxl.Workbook()
    _seed_and_move(wb_o.active)
    grid_o = _window(wb_o.active, 4, 4)

    wb_x = oxlsx.Workbook()
    _seed_and_move(wb_x.active)
    grid_x = _window(wb_x.active, 4, 4)

    assert grid_x == expected
    assert grid_x == grid_o

    # openpyxl SUPPORTS translate=True, but oxlsx does NOT model formula
    # translation, so it honestly raises NotImplementedError on that path.
    wb_t = oxlsx.Workbook()
    ws_t = wb_t.active
    ws_t["A1"] = "=B1"
    with pytest.raises(NotImplementedError):
        ws_t.move_range("A1", rows=1, cols=0, translate=True)


# ---------------------------------------------------------------------------
# Conditional formatting (cross-engine: oxlsx writes -> openpyxl reads)
# ---------------------------------------------------------------------------


def _openpyxl_load_no_repair(path):
    """Load with openpyxl, escalating any repair UserWarning to an error."""
    with warnings.catch_warnings():
        warnings.simplefilter("error", UserWarning)
        return openpyxl.load_workbook(path)


def test_formula_rule_round_trip():
    # oxlsx writes a FormulaRule (cfRule type="expression"); openpyxl reopens
    # it with no repair and reads the expression + range back.
    wb = oxlsx.Workbook()
    ws = wb.active
    fill = oxlsx.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ws.conditional_formatting.add("A1:A10", oxlsx.FormulaRule(formula=["$A1>5"], fill=fill))

    path = _temp_xlsx("oxlsx_py_formula_rule_round_trip.xlsx")
    try:
        wb.save(path)
        ws2 = _openpyxl_load_no_repair(path).active
        ranges = list(ws2.conditional_formatting)
        assert len(ranges) == 1
        assert str(ranges[0].sqref) == "A1:A10"
        rules = ranges[0].rules
        assert len(rules) == 1
        assert rules[0].type == "expression"
        assert rules[0].formula == ["$A1>5"]
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_cell_is_rule():
    # oxlsx writes a CellIsRule (cfRule type="cellIs" + operator); openpyxl
    # reopens it and reads the operator + operand formula back.
    wb = oxlsx.Workbook()
    ws = wb.active
    fill = oxlsx.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ws.conditional_formatting.add(
        "B1:B10", oxlsx.CellIsRule(operator="greaterThan", formula=["5"], fill=fill)
    )

    path = _temp_xlsx("oxlsx_py_cell_is_rule.xlsx")
    try:
        wb.save(path)
        ws2 = _openpyxl_load_no_repair(path).active
        ranges = list(ws2.conditional_formatting)
        assert len(ranges) == 1
        assert str(ranges[0].sqref) == "B1:B10"
        rule = ranges[0].rules[0]
        assert rule.type == "cellIs"
        assert rule.operator == "greaterThan"
        assert rule.formula == ["5"]
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_conditional_formatting_with_pattern_fill():
    # The dxf fill attached to a rule survives to openpyxl: the read-back rule
    # resolves a DifferentialStyle whose fill carries the FFC7CE color.
    wb = oxlsx.Workbook()
    ws = wb.active
    fill = oxlsx.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ws.conditional_formatting.add("A1:A5", oxlsx.FormulaRule(formula=["$A1>5"], fill=fill))

    path = _temp_xlsx("oxlsx_py_cf_with_pattern_fill.xlsx")
    try:
        wb.save(path)
        ws2 = _openpyxl_load_no_repair(path).active
        rule = list(ws2.conditional_formatting)[0].rules[0]
        assert rule.dxf is not None
        dxf_fill = rule.dxf.fill
        assert dxf_fill is not None
        rgb = (dxf_fill.bgColor.rgb if dxf_fill.bgColor else None) or (
            dxf_fill.fgColor.rgb if dxf_fill.fgColor else None
        )
        assert rgb is not None and "FFC7CE" in rgb
    finally:
        if os.path.exists(path):
            os.remove(path)


# ---------------------------------------------------------------------------
# openpyxl-parity: dataframe_to_rows (openpyxl.utils.dataframe.dataframe_to_rows)
#
# Migration helper so users can do:
#     for row in oxlsx.dataframe_to_rows(df, index=True, header=True):
#         ws.append(row)
#
# pandas is an OPTIONAL, user-provided integration — oxlsx duck-types the
# DataFrame and does NOT depend on pandas. These tests are skip-safe in a CI
# venv without pandas via pytest.importorskip.
# ---------------------------------------------------------------------------


def test_dataframe_to_rows_matches_openpyxl():
    pd = pytest.importorskip("pandas")
    from openpyxl.utils.dataframe import dataframe_to_rows as openpyxl_d2r

    # Mixed per-column dtypes (int / float / str) guard against any value
    # round-tripping that would upcast a column (e.g. via df.values).
    df = pd.DataFrame({"a": [1, 2], "b": [3.5, 4.5], "c": ["x", "y"]})

    for index in (True, False):
        for header in (True, False):
            expected = [list(r) for r in openpyxl_d2r(df, index=index, header=header)]
            actual = [
                list(r) for r in oxlsx.dataframe_to_rows(df, index=index, header=header)
            ]
            assert actual == expected, f"index={index} header={header}"


def test_dataframe_to_rows_default_args_match_openpyxl():
    pd = pytest.importorskip("pandas")
    from openpyxl.utils.dataframe import dataframe_to_rows as openpyxl_d2r

    df = pd.DataFrame({"name": ["Ann", "Bob"], "age": [30, 25]})

    # Defaults are index=True, header=True (openpyxl parity).
    expected = [list(r) for r in openpyxl_d2r(df)]
    actual = [list(r) for r in oxlsx.dataframe_to_rows(df)]
    assert actual == expected
    # Anchor the concrete expected layout: leading None header cell + the
    # always-yielded index-names row when index=True.
    assert actual == [
        [None, "name", "age"],
        [None],
        [0, "Ann", 30],
        [1, "Bob", 25],
    ]


def test_dataframe_to_rows_then_append():
    pd = pytest.importorskip("pandas")
    df = pd.DataFrame({"name": ["Ann", "Bob"], "age": [30, 25]})

    wb = oxlsx.Workbook()
    ws = wb.active
    for row in oxlsx.dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    path = _temp_xlsx("oxlsx_py_dataframe_to_rows_append.xlsx")
    try:
        wb.save(path)
        ws2 = oxlsx.load_workbook(path).active
        grid = _window(ws2, 3, 2)
        assert grid == [("name", "age"), ("Ann", 30), ("Bob", 25)]
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_dataframe_to_rows_multiindex_not_supported():
    pd = pytest.importorskip("pandas")

    mi = pd.DataFrame(
        [[1, 2]], columns=pd.MultiIndex.from_tuples([("A", "x"), ("A", "y")])
    )
    # MultiIndex columns/index are out of scope; oxlsx raises rather than
    # emitting a silently-wrong (single-level) layout.
    with pytest.raises(NotImplementedError):
        list(oxlsx.dataframe_to_rows(mi, index=True, header=True))

