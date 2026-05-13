#!/usr/bin/env python3
"""
Benchmark comparison: openpyxl vs oxlsx (Rust)
Measures the same operations on both sides for a fair comparison.

Usage:
    python3 benches/openpyxl_comparison.py

Requirements:
    pip install openpyxl
"""

import time
import tempfile
import os
import statistics
import openpyxl
from datetime import date, timedelta


SAMPLES_DIR = os.path.join(os.path.dirname(__file__), "..", "samples")
BENCH_10K    = os.path.join(SAMPLES_DIR, "bench_10k.xlsx")
ITERATIONS   = 5


def measure(label: str, fn, iterations: int = ITERATIONS):
    times = []
    for _ in range(iterations):
        t0 = time.perf_counter()
        fn()
        times.append(time.perf_counter() - t0)
    med = statistics.median(times) * 1000
    mn  = min(times) * 1000
    print(f"  {label:<45} median={med:>8.2f}ms  min={mn:>8.2f}ms")
    return med


print("=" * 70)
print("openpyxl benchmark (Python baseline)")
print(f"openpyxl version: {openpyxl.__version__}")
print("=" * 70)

# ── READ benchmarks ───────────────────────────────────────────────────────────
print("\n[READ]")

def read_10k_metadata():
    wb = openpyxl.load_workbook(BENCH_10K, read_only=True, data_only=True)
    _ = wb.sheetnames
    wb.close()

def read_10k_all_cells():
    wb = openpyxl.load_workbook(BENCH_10K, read_only=True, data_only=True)
    ws = wb.active
    count = 0
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            _ = cell
            count += 1
    wb.close()
    return count

measure("open bench_10k.xlsx (metadata only, read_only=True)", read_10k_metadata)
measure("open bench_10k.xlsx + iterate ALL cells (read_only)", read_10k_all_cells)


def read_10k_full():
    wb = openpyxl.load_workbook(BENCH_10K, data_only=True)
    ws = wb.active
    count = 0
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            _ = cell
            count += 1
    return count

measure("open bench_10k.xlsx + iterate ALL cells (full load)", read_10k_full)

# ── WRITE benchmarks ──────────────────────────────────────────────────────────
print("\n[WRITE]")

def write_single_cell():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Hello"
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    os.unlink(path)

def write_1k_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 1001):
        for col in range(1, 6):
            ws.cell(row=row, column=col, value=float(row))
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    os.unlink(path)

def write_10k_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 10001):
        for col in range(1, 6):
            ws.cell(row=row, column=col, value=float(row))
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    os.unlink(path)

def write_multi_sheet_3x1k():
    wb = openpyxl.Workbook()
    for i in range(3):
        ws = wb.active if i == 0 else wb.create_sheet(f"Sheet{i+1}")
        for row in range(1, 1001):
            ws.cell(row=row, column=1, value=f"row {row}")
            ws.cell(row=row, column=2, value=float(row))
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    os.unlink(path)

measure("write single cell",                 write_single_cell)
measure("write 1k rows × 5 cols (5k cells)", write_1k_rows)
measure("write 10k rows × 5 cols (50k cells)", write_10k_rows)
measure("write 3 sheets × 1k rows",          write_multi_sheet_3x1k)

print()
print("=" * 70)
print("Run Rust benchmarks with:")
print("  cargo bench")
print()
print("Then compare results manually.")
print("=" * 70)
